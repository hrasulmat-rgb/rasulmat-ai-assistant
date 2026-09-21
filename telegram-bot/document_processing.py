"""Safe local extraction helpers for Telegram document attachments."""

from dataclasses import dataclass
from datetime import date, datetime
import math
from pathlib import Path
from typing import Any


MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_TEXT_CHARS = 100_000
MAX_MEMORY_DOCUMENT_CHARS = 9_000

SUPPORTED_EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".txt": "text/plain",
    ".csv": "text/csv",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
}


class AttachmentError(Exception):
    """Base class for user-safe attachment errors."""


class UnsupportedAttachmentError(AttachmentError):
    """Raised when a file type is not supported."""


class AttachmentTooLargeError(AttachmentError):
    """Raised when a file exceeds the safe download limit."""


@dataclass(frozen=True)
class ProcessedAttachment:
    display_name: str
    extension: str
    mime_type: str
    extracted_text: str
    multimodal: bool


def extension_for(filename: str) -> str:
    """Return a normalized supported extension or raise a user-safe error."""
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedAttachmentError(
            f"Unsupported file type. Supported types: {supported}."
        )
    return extension


def validate_size(size: int | None) -> None:
    """Reject known oversized Telegram files before downloading them."""
    if size is not None and size > MAX_FILE_SIZE_BYTES:
        raise AttachmentTooLargeError(
            f"Files larger than {MAX_FILE_SIZE_BYTES // (1024 * 1024)} MB are not supported."
        )


def validate_downloaded_size(path: Path) -> None:
    """Reject a downloaded file if its actual size exceeds the limit."""
    size = path.stat().st_size
    if size > MAX_FILE_SIZE_BYTES:
        raise AttachmentTooLargeError(
            f"Files larger than {MAX_FILE_SIZE_BYTES // (1024 * 1024)} MB are not supported."
        )


def _clip_text(text: str, limit: int = MAX_EXTRACTED_TEXT_CHARS) -> str:
    text = text.strip()
    if len(text) <= limit:
        return text
    marker = "\n\n[Middle of extracted content omitted for safety.]\n\n"
    available = max(0, limit - len(marker))
    first_half = available // 2
    second_half = available - first_half
    return f"{text[:first_half]}{marker}{text[-second_half:]}"


def compact_for_memory(text: str, limit: int = MAX_MEMORY_DOCUMENT_CHARS) -> str:
    """Keep a bounded document excerpt for follow-up questions after cleanup."""
    return _clip_text(text, limit)


def _format_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        return format(value, ".15g")
    return str(value)


def _extract_text(path: Path) -> str:
    return _clip_text(path.read_text(encoding="utf-8", errors="replace"))


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    sections: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        rows = [f"Table {table_index}:"]
        for row in table.rows:
            rows.append("\t".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
        sections.append("\n".join(rows))
    return _clip_text("\n\n".join(sections))


def _xlsx_cell_text(formula_cell: Any, value_cell: Any) -> str:
    value = value_cell.value
    formula = formula_cell.value
    formatted_value = _format_value(value if value is not None else formula)
    if formula is not None and formula != value and str(formula).startswith("="):
        return f"{formatted_value} [formula: {formula}]"
    return formatted_value


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    formula_workbook = load_workbook(path, read_only=True, data_only=False)
    value_workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[str] = []
    try:
        for formula_sheet in formula_workbook.worksheets:
            value_sheet = value_workbook[formula_sheet.title]
            rows: list[str] = [f"Worksheet: {formula_sheet.title}"]
            value_rows = value_sheet.iter_rows()
            for formula_row in formula_sheet.iter_rows():
                value_row = next(value_rows, ())
                cells: list[str] = []
                for formula_cell, value_cell in zip(formula_row, value_row):
                    cell_text = _xlsx_cell_text(formula_cell, value_cell)
                    if cell_text:
                        cells.append(f"{formula_cell.coordinate}={cell_text}")
                if cells:
                    rows.append(" | ".join(cells))
            if len(rows) > 1:
                sections.append("\n".join(rows))
    finally:
        formula_workbook.close()
        value_workbook.close()
    return _clip_text("\n\n".join(sections))


def _extract_xls(path: Path) -> str:
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    sections: list[str] = []
    try:
        for sheet in workbook.sheets():
            rows: list[str] = [f"Worksheet: {sheet.name}"]
            for row_index in range(sheet.nrows):
                cells: list[str] = []
                for column_index in range(sheet.ncols):
                    cell = sheet.cell(row_index, column_index)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode)
                    cell_text = _format_value(value)
                    if cell_text:
                        coordinate = f"{xlrd.formula.colname(column_index)}{row_index + 1}"
                        cells.append(f"{coordinate}={cell_text}")
                if cells:
                    rows.append(" | ".join(cells))
            if len(rows) > 1:
                sections.append("\n".join(rows))
    finally:
        workbook.release_resources()
    return _clip_text("\n\n".join(sections))


def _extract_pdf(path: Path) -> str:
    import fitz

    pages: list[str] = []
    with fitz.open(path) as document:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append(f"Page {page_number}:\n{text}")
            else:
                pages.append(f"Page {page_number}: [No selectable text; inspect the scanned page image.]")
    return _clip_text("\n\n".join(pages))


def process_attachment(path: Path, display_name: str) -> ProcessedAttachment:
    """Extract local text and mark files that should also be sent to Gemini as media."""
    extension = extension_for(display_name)
    validate_downloaded_size(path)

    if extension in {".jpg", ".jpeg", ".png", ".pdf"}:
        extracted_text = _extract_pdf(path) if extension == ".pdf" else ""
        return ProcessedAttachment(
            display_name=display_name,
            extension=extension,
            mime_type=SUPPORTED_EXTENSIONS[extension],
            extracted_text=extracted_text,
            multimodal=True,
        )
    if extension in {".txt", ".csv"}:
        extracted_text = _extract_text(path)
    elif extension == ".docx":
        extracted_text = _extract_docx(path)
    elif extension == ".xlsx":
        extracted_text = _extract_xlsx(path)
    elif extension == ".xls":
        extracted_text = _extract_xls(path)
    else:
        raise UnsupportedAttachmentError(f"Unsupported file type: {extension}")

    return ProcessedAttachment(
        display_name=display_name,
        extension=extension,
        mime_type=SUPPORTED_EXTENSIONS[extension],
        extracted_text=extracted_text,
        multimodal=False,
    )